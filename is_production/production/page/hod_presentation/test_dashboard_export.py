from io import BytesIO
import unittest
from pptx import Presentation
from pptx.oxml.ns import qn
from is_production.production.report.hod_presentation.hod_presentation import build_hod_presentation


class TestDashboardExport(unittest.TestCase):
    def test_dashboard_capture_is_one_fitted_slide_without_a_duplicate_heading(self):
        pixel = 'iVBORw0KGgoAAAANSUhEUgAAAAoAAAAKCAIAAAACUFjqAAAAFUlEQVR4nGP8//8/A27AhEduBEsDAKXjAxF9kqZqAAAAAElFTkSuQmCC'
        output = BytesIO()
        build_hod_presentation({'site':'Uitgevallen','captured_slides':[
            {'title':'Uitgevallen · Production Performance','heading_style':'dashboard',
             'image_data':'data:image/png;base64,'+pixel}]}, output)
        output.seek(0)
        deck = Presentation(output)
        self.assertEqual(len(deck.slides), 1)
        slide = deck.slides[0]
        self.assertEqual(
            [s.text for s in slide.shapes if s.has_text_frame], ['1']
        )
        self.assertFalse(any(s.shape_type == 13 for s in slide.shapes))
        background = slide._element.cSld.bg
        self.assertIsNotNone(background)
        images = background.xpath('.//a:blip')
        self.assertEqual(len(images), 1)
        relationship = images[0].get(qn('r:embed'))
        self.assertTrue(relationship)
        self.assertTrue(
            slide.part.related_part(relationship).content_type.startswith('image/')
        )

    def test_two_sites_export_two_dashboard_slides(self):
        pixel = 'data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAoAAAAKCAIAAAACUFjqAAAAFUlEQVR4nGP8//8/A27AhEduBEsDAKXjAxF9kqZqAAAAAElFTkSuQmCC'
        captures = [
            {'title':f'{site} · Production Performance', 'heading_style':'dashboard', 'image_data':pixel}
            for site in ('Uitgevallen','Klipfontein')
        ]
        output = BytesIO()
        build_hod_presentation({'site':'Uitgevallen / Klipfontein','captured_slides':captures},output)
        output.seek(0)
        deck = Presentation(output)
        self.assertEqual(len(deck.slides),2)
        for slide in deck.slides:
            self.assertEqual([s.text for s in slide.shapes if s.has_text_frame], [str(deck.slides.index(slide) + 1)])
            self.assertFalse(any(shape.shape_type == 13 for shape in slide.shapes))
            background = slide._element.cSld.bg
            self.assertIsNotNone(background)
            images = background.xpath('.//a:blip')
            self.assertEqual(len(images), 1)
            relationship = images[0].get(qn('r:embed'))
            self.assertTrue(relationship)
            self.assertTrue(
                slide.part.related_part(relationship).content_type.startswith('image/')
            )
